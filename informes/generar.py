import html, re
import json, sys, os
# Uso: python3 informes/generar.py datos.json carpeta_salida   (formato de datos: informes/ejemplo.json)
D=json.load(open(sys.argv[1])); OUT=sys.argv[2]; os.makedirs(OUT, exist_ok=True)
e=lambda t: html.escape(str(t or ""))
def dig(t): return re.sub(r"\D","",t)
def card(p,i):
    c=[]
    if p["wa"]:
        href=p["wa"] if p["wa"].startswith("http") else "https://wa.me/"+dig(p["wa"])
        c.append(f'<a class="wa" href="{e(href)}">WhatsApp {e(p["wa"] if not p["wa"].startswith("http") else "")}</a>')
    if p["tel"]: c.append(f'<span>Tel. {e(p["tel"])}</span>')
    if p["mail"]: c.append(f'<a href="mailto:{e(p["mail"])}">{e(p["mail"])}</a>')
    host=re.sub(r"^https?://(www\.)?","",p["web"]).split("/")[0]
    return f'''<article class="card"><div class="n">{i}</div><div>
<div class="eti"><span>{e(p["tipo"])}</span><span>{e(p["calidad"])}</span></div>
<h3>{e(p["nombre"])}</h3><p class="gris">{e(p["lugar"])}</p>
<p><b>Marcas / productos:</b> {e(p["marcas"])}</p><p>{e(p["desc"])}</p>
<div class="dat"><div><small>Mínimo</small>{e(p["minimo"])}</div><div><small>Precios</small>{e(p["precio"])}</div></div>
{f'<p class="aviso">{e(p["nota"])}</p>' if p["nota"] else ""}
<p class="contactos">{"".join(c)}<a href="{e(p["web"])}">🌐 {e(host)}</a></p></div></article>'''
def seccion(titulo,lista,ini):
    out=f"<h2>{titulo}</h2>"; paises=[]
    for p in lista:
        if p["pais"] not in paises: paises.append(p["pais"])
    i=ini
    for pa in paises:
        out+=f'<h4>{pa}</h4>'
        for p in [x for x in lista if x["pais"]==pa]: out+=card(p,i); i+=1
    return out,i
cuerpo=""; i=1
for rub in D["rubros"]:
    h,i=seccion(rub["titulo"],rub["proveedores"],i); cuerpo+=h
TODOS=[(r["titulo"],p) for r in D["rubros"] for p in r["proveedores"]]
ref="".join(f"<tr><td>{e(a)}</td><td>{e(b)}</td></tr>" for a,b in D.get("referencias",[]))
doc=f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Informe de proveedores · {e(D["cliente"])} · Te Importamos</title><style>
@page{{size:A4;margin:13mm 12mm}}:root{{--n:#EA5B0C;--t:#1f2328;--g:#6b7280;--b:#e8e5e1;--f:#faf8f6}}
body{{margin:0;font-family:Helvetica,Arial,sans-serif;color:var(--t);font-size:12px;line-height:1.4}}
header{{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:3px solid var(--n);padding-bottom:12px;margin-bottom:16px}}
.marca{{font-size:24px;font-weight:800}}.marca i{{color:var(--n);font-style:normal}}header small{{color:var(--g);display:block;font-size:11px}}.fecha{{text-align:right;color:var(--g);font-size:11px}}
h1{{font-size:21px;margin:0 0 4px}}.ficha{{display:grid;grid-template-columns:repeat(4,1fr);gap:8px;margin:12px 0 14px}}
.ficha div{{background:var(--f);border:1px solid var(--b);border-radius:9px;padding:8px 10px}}.ficha small{{color:var(--g);display:block;font-size:10px;text-transform:uppercase}}.ficha b{{font-size:14px}}
.intro{{background:var(--f);border:1px solid var(--b);border-radius:10px;padding:10px 12px;margin-bottom:6px}}.intro ul{{margin:4px 0 0;padding-left:18px}}
h2{{font-size:15px;margin:22px 0 6px;padding:7px 10px;background:#fff3ec;border-left:4px solid var(--n);border-radius:4px;break-after:avoid}}
h4{{margin:12px 0 6px;font-size:13px;break-after:avoid}}
.card{{display:grid;grid-template-columns:26px 1fr;gap:10px;border:1px solid var(--b);border-radius:11px;padding:10px 12px;margin-bottom:8px;break-inside:avoid}}
.n{{width:24px;height:24px;border-radius:50%;background:var(--n);color:#fff;font-weight:800;display:flex;align-items:center;justify-content:center;font-size:11px}}
.card h3{{margin:3px 0 1px;font-size:14px}}.card p{{margin:3px 0}}.gris{{color:var(--g);font-size:11px}}
.eti{{display:flex;gap:5px;flex-wrap:wrap}}.eti span{{background:#fff3ec;color:#b2440a;border-radius:6px;padding:1px 7px;font-size:10px;font-weight:700}}
.dat{{display:grid;grid-template-columns:1fr 2fr;gap:8px;margin:5px 0}}.dat div{{background:var(--f);border-radius:7px;padding:5px 8px}}.dat small{{display:block;color:var(--g);font-size:10px;text-transform:uppercase}}
.aviso{{background:#fffbeb;color:#92400e;border-radius:6px;padding:4px 7px;font-size:11px}}
.contactos{{display:flex;gap:4px 12px;flex-wrap:wrap;font-weight:600;margin-top:5px}}.contactos a{{color:#1d4ed8;text-decoration:none}}.contactos a.wa{{color:#16a34a}}
table{{border-collapse:collapse;width:100%;margin-top:6px}}td{{border-bottom:1px solid var(--b);padding:5px 6px}}
.nota{{margin-top:18px;padding:10px 12px;border:1px solid var(--b);border-radius:10px;color:var(--g);font-size:10.5px;background:var(--f);break-inside:avoid}}
footer{{margin-top:14px;display:flex;justify-content:space-between;color:var(--g);font-size:10px;border-top:1px solid var(--b);padding-top:8px}}
</style></head><body>
<header><div><div class="marca">Te <i>Importamos</i></div><small>Importación por encargo · Rosario, Argentina · teimportamosarg.com</small></div><div class="fecha">Informe de búsqueda de proveedores<br>{e(D["fecha"])}</div></header>
<h1>{e(D["titulo"])}</h1><p class="gris">Preparado para: <b>{e(D["cliente"])}</b></p>
<div class="ficha"><div><small>Pedido</small><b>{e(D.get("cantidad") or "-")}</b></div><div><small>Calidad</small><b>{e(D["calidad"])}</b></div><div><small>Proveedores</small><b>{len(TODOS)}</b></div><div><small>Países</small><b>{e(D["paises"])}</b></div></div>
<div class="intro"><b>Lo más importante</b><ul>{"".join("<li>"+e(x)+"</li>" for x in D["claves"])}</ul></div>
{cuerpo}
{('<h2>💲 Precios de referencia</h2><table>'+ref+'</table>') if ref else ""}
<div class="nota"><b>Cómo leer este informe.</b> Todos los proveedores y datos de contacto fueron relevados en sus páginas web el {e(D["fecha"])}. Precios y stock cambian seguido: se confirman al consultar. Recomendaciones: pedir fotos reales y factura, empezar con una compra chica, y en celulares confirmar si vienen <b>liberados</b> (sin bloqueo de operadora) y su <b>grado de estado</b>. Te Importamos puede gestionar la compra, el envío y la importación completa.</div>
<footer><span>Te Importamos · WhatsApp 341 805-1515</span><span>Informe de proveedores · {e(D["cliente"])}</span></footer></body></html>'''
open(os.path.join(OUT,"informe.html"),"w").write(doc)
# Planilla
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb=Workbook(); ws=wb.active; ws.title="Proveedores"
cols=["#","Rubro","País","Proveedor","Ubicación","Tipo","Calidad","Marcas / productos","Mínimo","Precios","WhatsApp","Teléfono","Mail","Web","Notas"]
ws.append(cols)
for i,(rub,p) in enumerate(TODOS,1):
    ws.append([i,rub,re.sub(r"^\W+","",p["pais"]).strip(),p["nombre"],p["lugar"],p["tipo"],p["calidad"],p["marcas"],p["minimo"],p["precio"],p["wa"],p["tel"],p["mail"],p["web"],p["nota"]])
    ws.cell(i+1,14).hyperlink=p["web"]; ws.cell(i+1,14).font=Font(name="Arial",color="1D4ED8",underline="single")
    if p["wa"]:
        ws.cell(i+1,11).hyperlink=p["wa"] if p["wa"].startswith("http") else "https://wa.me/"+dig(p["wa"])
for row in ws.iter_rows():
    for cl in row:
        if cl.column!=14 or cl.row==1: cl.font=Font(name="Arial",bold=cl.row==1,color="FFFFFF" if cl.row==1 else "000000")
        cl.alignment=Alignment(wrap_text=True,vertical="top")
for cl in ws[1]: cl.fill=PatternFill("solid",fgColor="EA5B0C")
for col,w in zip("ABCDEFGHIJKLMNO",[4,12,14,24,22,20,20,30,18,34,18,16,24,30,30]): ws.column_dimensions[col].width=w
ws.freeze_panes="E2"; ws.auto_filter.ref=ws.dimensions
base="Proveedores-"+re.sub(r"[^A-Za-z0-9]+","-",D["cliente"]).strip("-")+"-TeImportamos"
wb.save(os.path.join(OUT,base+".xlsx")); open(os.path.join(OUT,"nombre.txt"),"w").write(base)
print(os.path.join(OUT,base))
